from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse, JsonResponse
from django.contrib import messages
from django.db.models import Q
from django.utils import timezone
from django.urls import reverse
from django.utils.text import slugify
from worklist.models import Study
from .models import Report
import io, base64
try:
    import fitz
except Exception:
    fitz = None
try:
    from docx import Document
except Exception:
    Document = None

@login_required
def report_list(request):
    if not getattr(request.user, 'can_edit_reports', None) or not request.user.can_edit_reports():
        return HttpResponse(status=403)
    search_query = request.GET.get('search', '')
    qs = Report.objects.select_related('study', 'study__patient', 'study__modality').all().order_by('-report_date')
    if search_query:
        qs = qs.filter(
            Q(study__patient__first_name__icontains=search_query) |
            Q(study__patient__last_name__icontains=search_query) |
            Q(study__accession_number__icontains=search_query)
        )
    context = {
        'reports': qs,
        'total_reports': Report.objects.count(),
        'draft_reports': Report.objects.filter(status='draft').count(),
        'pending_reports': Report.objects.filter(status='preliminary').count(),
        'final_reports': Report.objects.filter(status='final').count(),
        'search_query': search_query,
    }
    return render(request, 'reports/report_list.html', context)

@login_required
def write_report(request, study_id):
    if not getattr(request.user, 'can_edit_reports', None) or not request.user.can_edit_reports():
        return HttpResponse(status=403)
    study = get_object_or_404(Study, id=study_id)
    report = Report.objects.filter(study=study).first()
    is_new = report is None
    if request.method == 'POST':
        data = request.POST
        if is_new:
            report = Report.objects.create(
                study=study,
                radiologist=request.user,
                clinical_history=data.get('clinical_history',''),
                technique=data.get('technique',''),
                comparison=data.get('comparison',''),
                findings=data.get('findings',''),
                impression=data.get('impression',''),
                recommendations=data.get('recommendations',''),
                status=data.get('status','draft')
            )
            messages.success(request, 'Report created successfully')
        else:
            report.clinical_history = data.get('clinical_history','')
            report.technique = data.get('technique','')
            report.comparison = data.get('comparison','')
            report.findings = data.get('findings','')
            report.impression = data.get('impression','')
            report.recommendations = data.get('recommendations','')
            report.status = data.get('status','draft')
            report.last_modified = timezone.now()
            if data.get('status') == 'final' or data.get('action') == 'submit':
                report.signed_date = timezone.now()
            report.save()
            messages.success(request, 'Report updated successfully')
        if report.status == 'final' or data.get('action') == 'submit':
            study.status = 'completed'
            study.save(update_fields=['status'])
            return redirect('reports:report_list')
        return redirect('reports:write_report', study_id=study.id)
    return render(request, 'reports/write_report.html', {'study': study, 'report': report, 'is_new_report': is_new})


@login_required
def print_report_stub(request, study_id):
    study = get_object_or_404(Study, id=study_id)
    report = Report.objects.filter(study=study).first()
    viewer_url = request.build_absolute_uri(reverse('dicom_viewer:viewer')) + f"?study={study.id}"
    report_url = request.build_absolute_uri(reverse('reports:print_report', args=[study.id]))
    def _qr_b64(text: str) -> str:
        try:
            import qrcode
            from io import BytesIO
            buf = BytesIO()
            qrcode.make(text).save(buf, format='PNG')
            return 'data:image/png;base64,' + base64.b64encode(buf.getvalue()).decode('ascii')
        except Exception:
            return ''
    qr_viewer = _qr_b64(viewer_url)
    qr_report = _qr_b64(report_url)
    html = f"""
    <html><head><title>Report {study.accession_number}</title></head>
    <body>
      <h2>Report: {study.accession_number}</h2>
      <div>Patient: {study.patient.full_name} ({study.patient.patient_id})</div>
      <div>Modality: {study.modality.code} &nbsp; Date: {study.study_date}</div>
      <hr/>
      <div><b>Clinical History</b><br>{(report.clinical_history if report else (study.clinical_info or '')) or '-'}</div>
      <div><b>Technique</b><br>{(report.technique if report else '') or '-'}</div>
      <div><b>Comparison</b><br>{(report.comparison if report else '') or '-'}</div>
      {f'<div><b>Findings</b><br>{(report.findings or "").strip()}</div>' if report else ''}
      {f'<div><b>Impression</b><br>{(report.impression or "").strip()}</div>' if report else ''}
      <hr/>
      <div style="display:flex; gap:40px;">
        <div><img src="{qr_viewer}" height="100"/><div>View images</div></div>
        <div><img src="{qr_report}" height="100"/><div>View report</div></div>
      </div>
      <script>window.print && window.print()</script>
    </body></html>
    """
    return HttpResponse(html)


@login_required
def export_report_pdf(request, study_id):
    if not getattr(request.user, 'can_edit_reports', None) or not request.user.can_edit_reports():
        return HttpResponse(status=403)
    if fitz is None:
        return JsonResponse({'error': 'PDF export not available (PyMuPDF missing).'}, status=500)
    study = get_object_or_404(Study, id=study_id)
    report = Report.objects.filter(study=study).first()
    filename = f"report_{slugify(study.accession_number)}.pdf"
    doc = fitz.open(); page = doc.new_page(); margin = 36; y = margin
    for line in [
        f"Radiology Report",
        f"Patient: {study.patient.full_name} ({study.patient.patient_id})",
        f"Accession: {study.accession_number}    Modality: {study.modality.code}    Date: {study.study_date}",
    ]:
        page.insert_text((margin, y), line, fontsize=12 if line=="Radiology Report" else 10, fontname='helv', fill=(0,0,0)); y += 18
    def section(title, content):
        nonlocal y
        page.insert_text((margin, y), title, fontsize=11, fontname='helv', fill=(0,0,0)); y += 14
        page.insert_text((margin, y), (content or '-'), fontsize=10, fontname='helv', fill=(0,0,0)); y += 18
    section('Clinical History', (report.clinical_history if report else (study.clinical_info or '')))
    section('Technique', (report.technique if report else ''))
    section('Comparison', (report.comparison if report else ''))
    section('Findings', (report.findings if report else ''))
    section('Impression', (report.impression if report else ''))
    section('Recommendations', (report.recommendations if report else ''))
    buf = io.BytesIO(); doc.save(buf); buf.seek(0)
    resp = HttpResponse(buf.getvalue(), content_type='application/pdf')
    resp['Content-Disposition'] = f'attachment; filename="{filename}"'
    return resp


@login_required
def export_report_docx(request, study_id):
    # Allow radiologists/admins and superusers to export
    if not (request.user.is_superuser or (getattr(request.user, 'can_edit_reports', None) and request.user.can_edit_reports())):
        return HttpResponse(status=403)
    if Document is None:
        return JsonResponse({'error': 'DOCX export not available (python-docx missing).'}, status=500)
    study = get_object_or_404(Study, id=study_id)
    report = Report.objects.filter(study=study).first()
    doc = Document()
    # Optional letterhead if facility has one
    facility = study.facility
    try:
        if getattr(facility, 'letterhead', None) and getattr(facility.letterhead, 'url', None):
            from docx.shared import Inches
            from django.conf import settings as dj_settings
            import os as _os
            lh_path = _os.path.join(dj_settings.MEDIA_ROOT, facility.letterhead.name)
            if _os.path.exists(lh_path):
                doc.add_picture(lh_path, width=Inches(6))
    except Exception:
        pass
    doc.add_heading('Radiology Report', 0)
    doc.add_paragraph(f"Patient: {study.patient.full_name} ({study.patient.patient_id})")
    doc.add_paragraph(f"Accession: {study.accession_number}    Modality: {study.modality.code}    Date: {study.study_date}")
    for title, content in [
        ('Clinical History', (report.clinical_history if report else (study.clinical_info or ''))),
        ('Technique', (report.technique if report else '')),
        ('Comparison', (report.comparison if report else '')),
        ('Findings', (report.findings if report else '')),
        ('Impression', (report.impression if report else '')),
        ('Recommendations', (report.recommendations if report else '')),
    ]:
        doc.add_heading(title, level=2); doc.add_paragraph(content or '-')
    buf = io.BytesIO(); doc.save(buf); buf.seek(0)
    resp = HttpResponse(buf.getvalue(), content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    resp['Content-Disposition'] = f'attachment; filename="report_{slugify(study.accession_number)}.docx"'
    return resp
